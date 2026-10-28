#!/usr/bin/env python3
"""Import Taguig's figures from BLGF "by LGU" spreadsheets.

BLGF (blgf.gov.ph/lgu-fiscal-data) blocks scripted downloads, so download the
.xlsx files in a browser first. Then:

  # Statement of Receipts and Expenditures, one sheet per period
  python3 scripts/blgf/import_blgf.py sre ~/Downloads/SRE-QTR-2025.xlsx --sheet Q32025 --period Q3-2025

  # Full-year SRE report = cumulative Q4 of that year
  python3 scripts/blgf/import_blgf.py sre ~/Downloads/By-LGU-SRE-2025.xlsx --sheet "By LGU SRE FY2025" --period Q4-2025

  # Annual Regular Income and dependencies
  python3 scripts/blgf/import_blgf.py ari ~/Downloads/By-LGU-ARI-and-Dependencies-2025.xlsx --sheet "ARI 2025" --period FY-2025

Run without --sheet to list a workbook's sheets. Quarterly SRE figures are
cumulative year to date (Q2 = January to June), matching sre.json.
Values are converted from pesos to millions of pesos.
Requires: openpyxl
"""

import argparse
import json
import re
import sys
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parents[2]
SRE_PATH = ROOT / 'src' / 'data' / 'transparency' / 'sre.json'
ARI_PATH = ROOT / 'src' / 'data' / 'statistics' / 'ari.json'
LGU_NAME = 'TAGUIG CITY'
LOCATION = {'region': 'National Capital Region', 'province': 'Metro Manila',
            'lgu_name': 'Taguig', 'lgu_type': 'City'}

# Spot checks on the header text of key columns, so a layout change fails loudly.
SRE_HEADER_CHECKS = {5: 'CURRENT OPERATING INCOME', 17: 'National Tax Allotment',
                     22: 'TOTAL CURRENT OPERATING INCOME', 57: 'FUND/ CASH BALANCE, END'}
ARI_HEADER_CHECKS = {5: 'LOCALLY SOURCED REVENUE', 15: 'NATIONAL TAX ALLOTMENT',
                     23: 'ANNUAL REGULAR INCOME'}


def load_rows(path, sheet):
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    if sheet is None:
        sys.exit(f'Sheets in {path}: {wb.sheetnames}\nPass one with --sheet.')
    return list(wb[sheet].iter_rows(values_only=True))


def check_headers(rows, checks):
    header_text = {}
    for row in rows[5:12]:
        for i, cell in enumerate(row):
            if cell is not None:
                header_text.setdefault(i, []).append(re.sub(r'\s+', ' ', str(cell)))
    for col, expected in checks.items():
        found = ' '.join(header_text.get(col, []))
        if expected.lower() not in found.lower():
            sys.exit(f'Unexpected layout: column {col} is "{found}", expected "{expected}".')


def taguig_row(rows):
    matches = [r for r in rows if r[3] and str(r[3]).strip().upper() == LGU_NAME]
    if len(matches) != 1:
        sys.exit(f'Expected exactly one "{LGU_NAME}" row, found {len(matches)}.')
    return matches[0]


def millions(row, i):
    return round((row[i] or 0) / 1e6, 2)


def check_close(a, b, label):
    if abs(a - b) >= 2:  # pesos; allows for rounding in BLGF's own totals
        sys.exit(f'Totals do not add up ({label}): {a} vs {b}')


def build_sre(t, period):
    v = lambda i: t[i] or 0
    check_close(v(16) + v(21), v(22), 'local + external = operating income')
    check_close(v(10) + v(15), v(16), 'tax + non-tax = local sources')
    check_close(v(23) + v(29) + v(30) + v(31), v(32), 'operating expenditures')
    check_close(v(22) - v(32), v(33), 'net operating income')
    check_close(v(53) + v(52), v(54), 'fund/cash available')
    m = lambda i: millions(t, i)
    return {
        'period': period,
        'location_info': LOCATION,
        'current_operating_income': {
            'local_sources': {
                'tax_revenue': {
                    'real_property_tax': {'general_fund': m(5), 'special_education_fund': m(6), 'total': m(7)},
                    'tax_on_business': m(8), 'other_taxes': m(9), 'total_tax_revenue': m(10)},
                'non_tax_revenue': {
                    'regulatory_fees': m(11), 'service_user_charges': m(12),
                    'receipts_from_economic_enterprises': m(13), 'other_receipts': m(14),
                    'total_non_tax_revenue': m(15)},
                'total_local_sources': m(16)},
            'external_sources': {
                'national_tax_allotment': m(17), 'other_shares_from_national_tax_collection': m(18),
                'inter_local_transfers': m(19), 'extraordinary_receipts_grants_donations_aids': m(20),
                'total_external_sources': m(21)},
            'total_current_operating_income': m(22)},
        'total_current_operating_expenditures': {
            'general_public_services': m(23),
            'social_services': {
                'education_culture_sports_manpower_development': m(24),
                'health_nutrition_population_control': m(25), 'labor_and_employment': m(26),
                'housing_and_community_development': m(27),
                'social_services_and_social_welfare': m(28), 'total_social_services': m(29)},
            'economic_services': m(30), 'debt_service_interest_expense': m(31),
            'total_current_operating_expenditures': m(32)},
        'net_operating_income_loss_from_current_operations': m(33),
        'non_income_receipts': {
            'capital_investment_receipts': {
                'proceeds_from_sale_of_assets': m(34),
                'proceeds_from_sale_of_debt_securities_of_other_entities': m(35),
                'collection_of_loans_receivables': m(36), 'total_capital_investment_receipts': m(37)},
            'receipts_from_loans_and_borrowings': {
                'acquisition_of_loans': m(38), 'issuance_of_bonds': m(39),
                'total_receipts_from_loans_and_borrowings': m(40)},
            'other_non_income_receipts': m(41), 'total_non_income_receipts': m(42)},
        'non_operating_expenditures': {
            'capital_investment_expenditures': {
                'purchase_construct_of_property_plant_and_equipment': m(43),
                'purchase_of_debt_securities_of_other_entities': m(44),
                'grant_make_loan_to_other_entities': m(45), 'total_capital_investment_expenditures': m(46)},
            'debt_service_principal_cost': {
                'payment_of_loan_amortization': m(47), 'retirement_redemption_of_bonds_debt_securities': m(48),
                'total_debt_service_principal_cost': m(49)},
            'other_non_operating_expenditures': m(50), 'total_non_operating_expenditures': m(51)},
        'fund_summary': {
            'net_increase_decrease_in_funds': m(52), 'add_cash_balance_beginning': m(53),
            'fund_cash_available': m(54), 'less_payment_of_prior_years_accounts_payable': m(55),
            'continuing_appropriation': m(56), 'fund_cash_balance_end': m(57)},
    }


def build_ari(t, period):
    check_close(t[13] + t[14] + t[15] + t[22], t[23], 'LSR + interest + NTA + other shares = ARI')
    m = lambda i: millions(t, i)
    return {
        'period': period,
        'location_info': LOCATION,
        'locally_sourced_revenue': {
            'tax_revenue': {'real_property_tax_general_fund': m(5), 'tax_on_business': m(6),
                            'other_taxes': m(7), 'total_tax_revenue': m(8)},
            'non_tax_revenue': {'regulatory_fees': m(9), 'service_user_charges': m(10),
                                'receipts_from_economic_enterprises': m(11), 'total_non_tax_revenue': m(12)},
            'total_locally_sourced_revenue': m(13)},
        'other_income_sources': {'interest_income': m(14), 'national_tax_allotment': m(15)},
        'other_shares_from_national_tax_collection': {
            'share_from_economic_zone': m(16), 'share_from_evat': m(17), 'share_from_national_wealth': m(18),
            'share_from_pagcor_pcso_lotto': m(19), 'share_from_tobacco_excise_tax': m(20), 'others': m(21),
            'total_other_shares': m(22)},
        'summary_indicators': {
            'annual_regular_income': m(23),
            'dependency_rates': {'lsr_dependency': f'{t[24] * 100:.1f}%', 'nta_dependency': f'{t[25] * 100:.1f}%',
                                 'other_shares_from_national_tax_dependency': f'{t[26] * 100:.2f}%'}},
    }


def period_key(period):
    """'Q2-2025' -> (2025, 2); 'FY-2024' -> (2024, 5)."""
    q, year = period.split('-')
    return int(year), 5 if q == 'FY' else int(q.lstrip('Q'))


def upsert(path, entry):
    data = json.loads(path.read_text(encoding='utf-8')) if path.exists() else []
    data = [d for d in data if d['period'] != entry['period']] + [entry]
    data.sort(key=lambda d: period_key(d['period']))
    path.write_text(json.dumps(data, indent=2, ensure_ascii=False) + '\n', encoding='utf-8')
    return [d['period'] for d in data]


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument('kind', choices=['sre', 'ari'])
    ap.add_argument('xlsx')
    ap.add_argument('--sheet')
    ap.add_argument('--period', help='e.g. Q3-2025 (SRE) or FY-2025 (ARI)')
    args = ap.parse_args()

    rows = load_rows(args.xlsx, args.sheet)
    if not args.period or not re.fullmatch(r'(Q[1-4]|FY)-\d{4}', args.period):
        sys.exit('Pass --period like Q3-2025 or FY-2025.')
    if args.kind == 'sre':
        check_headers(rows, SRE_HEADER_CHECKS)
        periods = upsert(SRE_PATH, build_sre(taguig_row(rows), args.period))
        print(f'Updated {SRE_PATH.relative_to(ROOT)}: {periods}')
    else:
        check_headers(rows, ARI_HEADER_CHECKS)
        periods = upsert(ARI_PATH, build_ari(taguig_row(rows), args.period))
        print(f'Updated {ARI_PATH.relative_to(ROOT)}: {periods}')
    print('Run `npx prettier --write` on the file, then check the page in the browser.')


if __name__ == '__main__':
    main()
